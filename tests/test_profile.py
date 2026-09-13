from docx import Document


def test_profile_404_before_creation(client):
    resp = client.get("/api/profile")
    assert resp.status_code == 404


def test_profile_create_and_update(client):
    resp = client.post(
        "/api/profile",
        json={"full_name": "Test User", "target_titles": "ML Engineer", "remote_preference": "remote"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["full_name"] == "Test User"
    assert body["has_cv"] is False

    # Second POST updates the same row rather than creating a new one.
    resp = client.post("/api/profile", json={"full_name": "Updated Name", "target_titles": "ML Engineer"})
    assert resp.status_code == 200
    assert resp.json()["full_name"] == "Updated Name"
    assert resp.json()["id"] == body["id"]


def test_profile_rejects_bad_remote_preference(client):
    resp = client.post("/api/profile", json={"full_name": "X", "remote_preference": "sometimes"})
    assert resp.status_code == 422


def test_cv_upload_and_parsing(client, tmp_path):
    client.post("/api/profile", json={"full_name": "Test User"})

    doc = Document()
    doc.add_paragraph("Test User")
    doc.add_paragraph("Experienced in LangChain, RAG pipelines and LangGraph.")
    cv_path = tmp_path / "cv.docx"
    doc.save(str(cv_path))

    with open(cv_path, "rb") as f:
        resp = client.post(
            "/api/profile/cv",
            files={"file": ("cv.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert resp.status_code == 200
    assert resp.json()["has_cv"] is True
    assert resp.json()["cv_filename"] == "cv.docx"

    resp = client.get("/api/profile/cv/text")
    assert resp.status_code == 200
    assert "LangChain" in resp.json()["text"]


def test_cv_upload_rejects_unsupported_extension(client, tmp_path):
    client.post("/api/profile", json={"full_name": "Test User"})
    bad_file = tmp_path / "cv.xyz"
    bad_file.write_text("not a real cv")
    with open(bad_file, "rb") as f:
        resp = client.post("/api/profile/cv", files={"file": ("cv.xyz", f, "application/octet-stream")})
    assert resp.status_code == 400
