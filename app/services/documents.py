"""
Turns the LLM's tailored-CV JSON / cover-letter text into actual .docx
files on disk, and records them in the database.
"""
from datetime import datetime
from app.utils.time import utcnow
from pathlib import Path

from docx import Document
from docx.shared import Pt
from sqlmodel import Session

from app.models import DocType, GeneratedDocument, JobListing, Profile
from app.services import tailoring


def _docs_dir() -> Path:
    from app.config import DOCS_DIR
    return DOCS_DIR


def _write_cv_docx(path: Path, profile: Profile, tailored: dict) -> None:
    doc = Document()

    title = doc.add_heading(profile.full_name or "Curriculum Vitae", level=0)
    contact_bits = [b for b in [profile.email, profile.phone, profile.location] if b]
    if contact_bits:
        p = doc.add_paragraph(" | ".join(contact_bits))
        p.runs[0].font.size = Pt(10)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(tailored.get("summary", ""))

    for section in tailored.get("sections", []):
        doc.add_heading(section.get("heading", ""), level=1)
        for line in str(section.get("content", "")).split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith(("-", "*", "\u2022")):
                doc.add_paragraph(line.lstrip("-*\u2022 ").strip(), style="List Bullet")
            else:
                doc.add_paragraph(line)

    doc.save(str(path))


def _write_cover_letter_docx(path: Path, profile: Profile, job: JobListing, letter_text: str) -> None:
    doc = Document()
    doc.add_paragraph(utcnow().strftime("%d %B %Y"))
    doc.add_paragraph("")
    doc.add_paragraph(f"Re: {job.title} at {job.company}")
    doc.add_paragraph("")
    for para in letter_text.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    doc.add_paragraph("")
    doc.add_paragraph(profile.full_name or "")
    doc.save(str(path))


def create_tailored_documents(
    session: Session, job: JobListing, profile: Profile, generate_cover_letter: bool = True
) -> list[GeneratedDocument]:
    docs_dir = _docs_dir()
    docs_dir.mkdir(parents=True, exist_ok=True)
    created: list[GeneratedDocument] = []

    tailored_cv = tailoring.generate_tailored_cv(job, profile)
    cv_path = docs_dir / f"job{job.id}_cv_{utcnow().strftime('%Y%m%d%H%M%S')}.docx"
    _write_cv_docx(cv_path, profile, tailored_cv)

    cv_summary_text = tailored_cv.get("summary", "") + "\n\n" + tailored_cv.get("change_notes", "")
    cv_record = GeneratedDocument(
        job_id=job.id,
        doc_type=DocType.cv,
        content_text=cv_summary_text,
        file_path=str(cv_path.relative_to(docs_dir.parent)),
    )
    session.add(cv_record)
    created.append(cv_record)

    if generate_cover_letter:
        letter_text = tailoring.generate_cover_letter(job, profile)
        letter_path = docs_dir / f"job{job.id}_cover_letter_{utcnow().strftime('%Y%m%d%H%M%S')}.docx"
        _write_cover_letter_docx(letter_path, profile, job, letter_text)

        letter_record = GeneratedDocument(
            job_id=job.id,
            doc_type=DocType.cover_letter,
            content_text=letter_text,
            file_path=str(letter_path.relative_to(docs_dir.parent)),
        )
        session.add(letter_record)
        created.append(letter_record)

    session.commit()
    for record in created:
        session.refresh(record)
    return created
